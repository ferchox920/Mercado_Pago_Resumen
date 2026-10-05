import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

import mercadopago_pdf_a_excel as app


class ExtractionTests(unittest.TestCase):
    def test_movement_amount_is_not_confused_with_balance(self):
        for amount in ("$ 1.234,56", "$1234,56", "$\n1234,56", "$ 1234567,89"):
            with self.subTest(amount=amount):
                block = (
                    "01-10-2026\nTransferencia recibida Persona Ejemplo\n"
                    f"12345678901\n{amount}\n$ 9.000,00"
                )
                movement = app._parse_movement_block(block, "Ejemplo", "Octubre")
                self.assertIsNotNone(movement)
                self.assertEqual(movement.valor, app.parse_money(amount))

    def test_entradas_total_accepts_spacing_and_amount_variants(self):
        for amount in ("$ 1.234,56", "$1.234,56", "$\n1.234,56", "$ 1234,56"):
            with self.subTest(amount=amount):
                self.assertEqual(app.extract_entradas_total(f"Entradas:\n{amount}"), 1234.56)
        self.assertIsNone(app.extract_entradas_total("Sin total de entradas"))

    def test_pdf_to_excel_and_total_warning(self):
        import fitz
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as folder:
            pdf = Path(folder) / "ejemplo.pdf"
            output = Path(folder) / "ejemplo.xlsx"
            text = "\n".join([
                "Resumen de cuenta en pesos", "Titular Ejemplo",
                "CVU: 0000000000000000000000", "Periodo: Octubre 2026",
                "Entradas: $1250,50", "Detalle de movimientos",
                "01-10-2026", "Transferencia recibida Persona Ejemplo",
                "12345678901", "$1000,00", "$9000,00",
                "02-10-2026", "Liquidación de dinero",
                "12345678902", "$250,50", "$9250,50",
            ])
            with fitz.open() as doc:
                page = doc.new_page()
                page.insert_text((50, 50), text)
                doc.save(pdf)
            movements, warnings = app.process_pdf(pdf)
            self.assertEqual(len(movements), 2)
            self.assertEqual(warnings, [])
            app.build_excel(movements, output)
            wb = load_workbook(output)
            try:
                self.assertEqual(len(wb.sheetnames), 4)
                self.assertEqual(wb["Movimientos"]["E5"].value, 1000)
                self.assertEqual(wb["Resumen"]["C4"].value, 1250.50)
            finally:
                wb.close()
            with patch.object(app, "extract_text_from_pdf", return_value=text.replace(
                "Entradas: $1250,50", "Entradas: $\n1300,50"
            )):
                _, warnings = app.process_pdf(pdf)
            self.assertEqual(len(warnings), 1)
            self.assertIn("diferencia -50.00", warnings[0])


class GuiProcessingTests(unittest.TestCase):
    def make_app(self):
        gui = app.MercadoPagoExtractorApp.__new__(app.MercadoPagoExtractorApp)
        gui.pdf_paths = [Path("primero.pdf"), Path("segundo.pdf")]
        gui.output_folder = Path("destino_original")
        gui.root = Mock()
        gui.root.after.side_effect = lambda delay, callback, *args: callback(*args)
        for name in ("generate_button", "select_files_button", "select_folder_button",
                     "clear_button", "progress", "log", "log_message", "update_progress"):
            setattr(gui, name, Mock())
        return gui

    def test_generation_captures_selection_and_disables_buttons(self):
        gui = self.make_app()
        with patch.object(app.threading, "Thread") as thread:
            gui.generate_excels()
        self.assertEqual(thread.call_args.kwargs["args"], (
            (Path("primero.pdf"), Path("segundo.pdf")), Path("destino_original")
        ))
        thread.return_value.start.assert_called_once()
        for button in (gui.generate_button, gui.select_files_button,
                       gui.select_folder_button, gui.clear_button):
            button.configure.assert_called_once_with(state="disabled")

    def test_worker_keeps_original_destination_and_restores_buttons_after_error(self):
        gui = self.make_app()
        pdf_paths = tuple(gui.pdf_paths)
        output_folder = gui.output_folder

        def process(pdf):
            gui.pdf_paths.clear()
            gui.output_folder = None
            if pdf == pdf_paths[0]:
                raise app.MercadoPagoPdfError("Error de prueba")
            return [Mock()], []

        with patch.object(app, "process_pdf", side_effect=process) as parse, \
                patch.object(app, "get_safe_output_path", return_value=Path("resultado.xlsx")) as path, \
                patch.object(app, "build_excel") as build:
            gui.generate_excels_from_gui(pdf_paths, output_folder)
        self.assertEqual(parse.call_count, 2)
        path.assert_called_once_with(pdf_paths[1], output_folder)
        build.assert_called_once()
        gui.update_progress.assert_called_with(2, 2)
        for button in (gui.generate_button, gui.select_files_button,
                       gui.select_folder_button, gui.clear_button):
            button.configure.assert_called_once_with(state="normal")


if __name__ == "__main__":
    unittest.main()
